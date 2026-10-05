import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
F="Arial"
df=pd.read_csv("./diamonds_clean.csv")
n=len(df); L=n+1
wb=Workbook(); ws=wb.active; ws.title="Summary"; d=wb.create_sheet("Data")
d.append(list(df.columns)+["price_per_carat","ln_price","ln_carat"])
for i,r in enumerate(df.itertuples(index=False),start=2):
    d.append(list(r)+[f"=H{i}/B{i}",f"=LN(H{i})",f"=LN(B{i})"])
for c in d[1]: c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F3864")
d.freeze_panes="A2"; d.auto_filter.ref=f"A1:N{L}"
for col in "ABCDEFGHIJKLMN": d.column_dimensions[col].width=14
R=lambda c:f"Data!${c}$2:${c}${L}"
H=Font(name=F,bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F3864")
B=Font(name=F,bold=True); N=Font(name=F); IN=Font(name=F,color="0000FF")
def head(row,labels):
    for j,t in enumerate(labels,1):
        c=ws.cell(row,j,t); c.font=H; c.fill=HF; c.alignment=Alignment(horizontal="center",wrap_text=True)
def put(r,c,v,fmt=None,font=N):
    x=ws.cell(r,c,v); x.font=font
    if fmt: x.number_format=fmt
    return x
ws["A1"]="Diamonds price analysis"; ws["A1"].font=Font(name=F,bold=True,size=14)
ws["A2"]="Source: Seaborn diamonds.csv, cleaned (duplicates and impossible dimensions removed). All tables below are live formulas on the Data sheet."; ws["A2"].font=Font(name=F,italic=True,size=9)
# KPIs
head(4,["Metric","Value"])
k=[("Diamonds (rows)",f"=COUNT({R('A')})","#,##0"),("Mean price ($)",f"=AVERAGE({R('H')})","$#,##0"),
   ("Median price ($)",f"=MEDIAN({R('H')})","$#,##0"),("Mean carat",f"=AVERAGE({R('B')})","0.00"),
   ("Price skew (raw)",f"=SKEW({R('H')})","0.00"),("Price skew (log)",f"=SKEW({R('M')})","0.00"),
   ("Carat elasticity of price (log-log slope)",f"=SLOPE({R('M')},{R('N')})","0.000"),
   ("R-squared, log price vs log carat",f"=RSQ({R('M')},{R('N')})","0.000")]
for i,(a,f,fm) in enumerate(k,5): put(i,1,a); put(i,2,f,fm)
# Clarity
r0=15; ws.cell(r0-1,1,"Clarity effect with size held roughly constant (0.9 to 1.1 ct)").font=B
head(r0,["Clarity (low to high)","Diamonds in band","Mean price in band ($)","Mean price, all sizes ($)"])
cl=["I1","SI2","SI1","VS2","VS1","VVS2","VVS1","IF"]
for i,c in enumerate(cl,r0+1):
    put(i,1,c)
    put(i,2,f'=COUNTIFS({R("E")},A{i},{R("B")},">=0.9",{R("B")},"<=1.1")',"#,##0")
    put(i,3,f'=AVERAGEIFS({R("H")},{R("E")},A{i},{R("B")},">=0.9",{R("B")},"<=1.1")',"$#,##0")
    put(i,4,f'=AVERAGEIFS({R("H")},{R("E")},A{i})',"$#,##0")
ce=r0+len(cl)
# Cut
c0=ce+3; ws.cell(c0-1,1,"Cut paradox: worse cuts look pricier on average because the stones are bigger").font=B
head(c0,["Cut (low to high)","Diamonds","Mean carat","Mean price ($)","Mean price per carat ($)"])
for i,c in enumerate(["Fair","Good","Very Good","Premium","Ideal"],c0+1):
    put(i,1,c); put(i,2,f'=COUNTIFS({R("C")},A{i})',"#,##0"); put(i,3,f'=AVERAGEIFS({R("B")},{R("C")},A{i})',"0.00")
    put(i,4,f'=AVERAGEIFS({R("H")},{R("C")},A{i})',"$#,##0"); put(i,5,f'=AVERAGEIFS({R("L")},{R("C")},A{i})',"$#,##0")
# Models
m0=c0+9; ws.cell(m0-1,1,"Model comparison (5-fold CV, from Python; hardcoded inputs in blue)").font=B
head(m0,["Model","R-squared (log price)","MAPE on price"])
for i,(a,r,m) in enumerate([("Mean baseline",0.000,1.1026),("Carat only (OLS)",0.933,0.2090),("OLS log-log + grades",0.983,0.1043),("Gradient boosting",0.991,0.0744)],m0+1):
    put(i,1,a); put(i,2,r,"0.000",IN); put(i,3,m,"0.0%",IN)
ws.cell(m0+5,1,"Assumption: out-of-fold scores from diamonds_analysis.py, random 5-fold split, seed 0. Boosting beats OLS by 3.0 MAPE points (bootstrap 95% CI 2.93 to 3.06).").font=Font(name=F,italic=True,size=9)
ws.column_dimensions["A"].width=44
for col in "BCDE": ws.column_dimensions[col].width=22
ch=BarChart(); ch.type="col"; ch.title="Mean price by clarity, 0.9-1.1 ct"; ch.y_axis.title="Price ($)"; ch.x_axis.title="Clarity"
ch.add_data(Reference(ws,min_col=3,min_row=r0,max_row=ce),titles_from_data=True); ch.set_categories(Reference(ws,min_col=1,min_row=r0+1,max_row=ce))
ch.legend=None; ch.height=8; ch.width=16; ch.x_axis.delete=False; ch.y_axis.delete=False
ws.add_chart(ch,"G4")
wb.save("./diamonds_analysis.xlsx")
